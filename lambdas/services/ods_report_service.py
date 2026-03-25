import os
import tempfile
from datetime import datetime

from openpyxl.workbook import Workbook
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from enums.dynamo_filter import AttributeOperator
from enums.file_type import FileType
from enums.lambda_error import LambdaError
from enums.metadata_field_names import DocumentReferenceMetadataFields
from enums.patient_ods_inactive_status import PatientOdsInactiveStatus
from enums.report_type import ReportType
from enums.repository_role import RepositoryRole
from models.document_review import DocumentUploadReviewReference
from services.base.dynamo_service import DynamoDBService
from services.base.s3_service import S3Service
from services.search_document_review_service import DocumentUploadReviewService
from utils.audit_logging_setup import LoggingService
from utils.common_query_filters import (
    FinalStatusAndNotSuperseded,
    get_not_deleted_filter,
)
from utils.dynamo_query_filter_builder import DynamoQueryFilterBuilder
from utils.lambda_exceptions import OdsReportException
from utils.request_context import request_context

logger = LoggingService(__name__)


class OdsReportService:
    def __init__(self):
        self.dynamo_service = DynamoDBService()
        self.table_name = os.getenv("LLOYD_GEORGE_DYNAMODB_NAME")
        self.reports_bucket = os.getenv("STATISTICAL_REPORTS_BUCKET")
        self.temp_output_dir = ""
        download_report_aws_role_arn = os.getenv("PRESIGNED_ASSUME_ROLE")
        self.s3_service = S3Service(custom_aws_role=download_report_aws_role_arn)
        self.document_upload_review_service = DocumentUploadReviewService()

    def get_nhs_numbers_by_ods(
        self,
        ods_code: str,
        is_pre_signed_needed: bool = False,
        is_upload_to_s3_needed: bool = False,
        file_type_output: FileType = FileType.CSV,
    ):
        results = self.query_table_by_index(ods_code)
        nhs_numbers = {
            item.get(DocumentReferenceMetadataFields.NHS_NUMBER.value)
            for item in results
            if item.get(DocumentReferenceMetadataFields.NHS_NUMBER.value)
        }
        if is_upload_to_s3_needed:
            self.temp_output_dir = tempfile.mkdtemp()
        return self.create_and_save_ods_report(
            ods_code,
            nhs_numbers,
            is_pre_signed_needed,
            is_upload_to_s3_needed,
            file_type_output,
        )

    def get_documents_for_review(
        self,
        ods_code: str,
        output_file_type: FileType = FileType.CSV,
    ):
        if output_file_type != FileType.CSV:
            raise OdsReportException(400, LambdaError.UnsupportedFileType)

        query_filter = self.document_upload_review_service.build_review_dynamo_filter()

        results = self.document_upload_review_service.fetch_documents_from_table(
            search_key="Custodian",
            search_condition=ods_code,
            index_name="CustodianIndex",
            query_filter=query_filter,
        )

        self.temp_output_dir = tempfile.mkdtemp()

        return self.create_and_save_ods_report(
            ods_code=ods_code,
            data=results,
            create_pre_signed_url=True,
            upload_to_s3=True,
            report_type=ReportType.REVIEW,
        )

    def scan_table_with_filter(self, ods_code: str):
        ods_codes = [ods_code]
        authorization_user = getattr(request_context, "authorization", {})
        if (
            authorization_user
            and authorization_user.get("repository_role") == RepositoryRole.PCSE.value
        ):
            ods_codes = [
                PatientOdsInactiveStatus.SUSPENDED,
                PatientOdsInactiveStatus.DECEASED,
            ]
        ods_filter_expression = self.build_filter_expression(ods_codes)
        field_names_expression = ",".join(DocumentReferenceMetadataFields.list())

        results = []
        response = self.dynamo_service.scan_whole_table(
            table_name=self.table_name,
            project_expression=field_names_expression,
            filter_expression=ods_filter_expression,
        )
        results.extend(response)
        if not results:
            logger.info("No records found for ODS code {}".format(ods_code))
            raise OdsReportException(404, LambdaError.NoDataFound)

        return results

    @staticmethod
    def build_filter_expression(ods_codes: list[str]):
        filter_builder = DynamoQueryFilterBuilder()
        delete_filter_expression = get_not_deleted_filter(filter_builder)
        filter_builder.add_condition("CurrentGpOds", AttributeOperator.IN, ods_codes)
        ods_code_filter_expression = filter_builder.build()
        return delete_filter_expression & ods_code_filter_expression

    def query_table_by_index(self, ods_code: str):
        ods_codes = [ods_code]
        authorization_user = getattr(request_context, "authorization", {})
        if (
            authorization_user
            and authorization_user.get("repository_role") == RepositoryRole.PCSE.value
        ):
            ods_codes = [
                PatientOdsInactiveStatus.SUSPENDED,
                PatientOdsInactiveStatus.DECEASED,
            ]
        results = []
        for ods_code in ods_codes:
            response = self.dynamo_service.query_table(
                table_name=self.table_name,
                index_name="OdsCodeIndex",
                search_key=DocumentReferenceMetadataFields.CURRENT_GP_ODS.value,
                search_condition=ods_code,
                query_filter=FinalStatusAndNotSuperseded,
            )
            results += response

        if not results:
            logger.info("No records found for ODS code {}".format(ods_code))
            raise OdsReportException(404, LambdaError.NoDataFound)
        return results

    def create_and_save_ods_report(
        self,
        ods_code: str,
        data: set[str] | list[DocumentUploadReviewReference],
        create_pre_signed_url: bool = False,
        upload_to_s3: bool = False,
        file_type_output: FileType = FileType.CSV,
        report_type: ReportType = ReportType.PATIENT,
    ):
        file_name, local_file_path = self.create_ods_report(
            ods_code,
            data,
            file_type_output,
            report_type,
        )

        if upload_to_s3:
            self.save_report_to_s3(ods_code, file_name, local_file_path)

            if create_pre_signed_url:
                return self.get_pre_signed_url(ods_code, file_name)

    def create_ods_report(
        self,
        ods_code: str,
        data: set[str] | list[DocumentUploadReviewReference],
        file_type_output: FileType = FileType.CSV,
        report_type: ReportType = ReportType.PATIENT,
    ):
        file_name = self.get_file_name_for_report_type(
            report_type,
            ods_code,
            len(data),
            file_type_output,
        )

        local_file_path = os.path.join(self.temp_output_dir, file_name)
        match file_type_output:
            case FileType.CSV:
                if report_type == ReportType.PATIENT:
                    self.create_csv_report(local_file_path, data, ods_code)
                elif report_type == ReportType.REVIEW:
                    self.create_review_csv_report(local_file_path, data)
            case FileType.XLSX:
                self.create_xlsx_report(local_file_path, data, ods_code)
            case FileType.PDF:
                self.create_pdf_report(local_file_path, data, ods_code)
            case _:
                raise OdsReportException(400, LambdaError.UnsupportedFileType)
        logger.info(f"Query completed. {len(data)} items written to {file_name}.")

        return (file_name, local_file_path)

    def get_file_name_for_report_type(
        self,
        report_type: ReportType,
        ods_code: str,
        data_length: int,
        file_type_output: FileType,
    ):
        formatted_time = datetime.now().strftime("%Y-%m-%d_%H-%M")

        file_name_prepend = ""

        match report_type:
            case ReportType.PATIENT:
                file_name_prepend = "LloydGeorgeSummary_"
            case ReportType.REVIEW:
                file_name_prepend = "ReviewReport_"

        file_name = (
            file_name_prepend
            + ods_code
            + f"_{data_length}_"
            + formatted_time
            + "."
            + file_type_output
        )

        return file_name

    def create_csv_report(self, file_name: str, nhs_numbers: set[str], ods_code: str):
        with open(file_name, "w") as f:
            f.write(
                f"Total number of patients for ODS code {ods_code}: {len(nhs_numbers)}\n",
            )
            f.write("NHS Numbers:\n")
            f.writelines(f"{nhs_number}\n" for nhs_number in nhs_numbers)

    def create_review_csv_report(
        self,
        file_name: str,
        data: list[DocumentUploadReviewReference],
    ):
        headers = [
            "nhs_number",
            "review_reason",
            "document_snomed_code_type",
            "author",
            "upload_date",
        ]

        with open(file_name, "w") as f:
            full_line = ""
            for header in headers:
                full_line += f"{header},"
            full_line = full_line[:-1]  # remove the trailing comma
            f.write(f"{full_line}\n")

            for line in data:
                upload_date = datetime.fromtimestamp(line.upload_date).isoformat()
                f.write(
                    line.nhs_number
                    + ","
                    + line.review_reason
                    + ","
                    + line.document_snomed_code_type
                    + ","
                    + line.author
                    + ","
                    + upload_date
                    + "\n",
                )

    def create_xlsx_report(self, file_name: str, nhs_numbers: set[str], ods_code: str):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = (
            f"Total number of patients for ODS code {ods_code}: {len(nhs_numbers)}\n"
        )
        ws["A2"] = "NHS Numbers:\n"
        for row in nhs_numbers:
            ws.append([row])

        wb.save(file_name)

    def create_pdf_report(self, file_name: str, nhs_numbers: set[str], ods_code: str):
        c = canvas.Canvas(file_name, pagesize=letter)
        _, height = letter
        c.setFont("Helvetica-Bold", 16)
        x = 100
        y = 700
        c.drawString(x, height - 50, f"NHS numbers within NDR for ODS code: {ods_code}")
        c.setFont("Helvetica", 12)

        c.drawString(x, y, f"Total number of patients: {len(nhs_numbers)}")
        y -= 20
        c.drawString(x, y, "NHS Numbers:")
        y -= 20
        for row in nhs_numbers:
            if y < 40:
                c.showPage()
                y = height - 50

            c.drawString(100, y, row)
            y -= 20

        c.save()

    def save_report_to_s3(self, ods_code: str, file_name: str, temp_file_path: str):
        logger.info("Uploading the csv report to S3 bucket...")
        today = datetime.now().date()
        self.s3_service.upload_file(
            s3_bucket_name=self.reports_bucket,
            file_key=f"ods-reports/{ods_code}/{today.year}/{today.month}/{today.day}/{file_name}",
            file_name=temp_file_path,
        )

    def get_pre_signed_url(self, ods_code: str, file_name: str):
        today = datetime.now().date()
        return self.s3_service.create_download_presigned_url(
            s3_bucket_name=self.reports_bucket,
            file_key=f"ods-reports/{ods_code}/{today.year}/{today.month}/{today.day}/{file_name}",
        )
